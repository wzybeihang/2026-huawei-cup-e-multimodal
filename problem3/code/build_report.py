"""Build competition-format report for Q2-model-on-Q3 experiment.
Uses OMML equations (convertible to MathType in Word) and Times New Roman font.
"""
import os
import sys
import json
import csv
import re
from pathlib import Path

from latex2mathml.converter import convert
from lxml import etree
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
R = ROOT / '实验结果'
F = ROOT / '图表'
OUT = ROOT / '报告'
OUT.mkdir(exist_ok=True)
FORM = OUT / '公式源码'
FORM.mkdir(exist_ok=True)


def j(p):
    return json.loads(Path(p).read_text(encoding='utf-8'))


def rows(p):
    with Path(p).open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


# ---- Load data ----
met = j(R / 'valid_metrics.json')
diag = j(R / 'report_diagnostics.json')
audit = j(R / 'attribution_audit.json')
pred = rows(ROOT / '提交结果/问题3_附件4预测与解释结果_Q2模型.csv')
faith = rows(R / '解释忠实性删除实验.csv')

# ---- Document setup ----
doc = Document()
s = doc.sections[0]
s.page_width = Cm(21)
s.page_height = Cm(29.7)
s.top_margin = Cm(2.5)
s.bottom_margin = Cm(2.5)
s.left_margin = Cm(3.0)
s.right_margin = Cm(2.5)

for style in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
    st = doc.styles[style]
    st.font.name = 'Times New Roman'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')

doc.styles['Normal'].font.size = Pt(12)
doc.styles['Normal'].paragraph_format.line_spacing = 1.5
doc.styles['Normal'].paragraph_format.space_after = Pt(0)
doc.styles['Normal'].paragraph_format.first_line_indent = Cm(0.74)

for name, size in [('Title', 16), ('Heading 1', 14), ('Heading 2', 13), ('Heading 3', 12)]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.font.bold = True
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '黑体')
    st.paragraph_format.space_before = Pt(12)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.first_line_indent = Cm(0)

doc.styles['Caption'].font.size = Pt(10.5)
doc.styles['Caption'].paragraph_format.first_line_indent = Cm(0)

doc.core_properties.title = '问题3 可解释性多模态情感预测——问题2方法迁移验证'

# ---- Equation helper ----
xsl_path = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'
xsl = etree.XSLT(etree.parse(xsl_path))
equations = []
prose = []


def p(t, indent=True):
    para = doc.add_paragraph(t)
    if not indent:
        para.paragraph_format.first_line_indent = Cm(0)
    prose.append(t)


def h(t, level=1):
    doc.add_heading(t, level=level)
    prose.append('\n' + '#' * level + ' ' + t)


def eq(tex, meaning):
    n = len(equations) + 1
    ml = etree.fromstring(convert(tex).encode())
    for empty in list(ml.findall('.//{http://www.w3.org/1998/Math/MathML}mo')):
        if not empty.text and not len(empty):
            empty.getparent().remove(empty)
    node = xsl(ml).getroot()
    if node.tag != qn('m:oMath'):
        node = node.find('.//' + qn('m:oMath'))
    assert node is not None and 'merror' not in etree.tostring(ml, encoding='unicode')
    for run in node.iter(qn('m:r')):
        pr = OxmlElement('w:rPr')
        fonts = OxmlElement('w:rFonts')
        fonts.set(qn('w:ascii'), 'Cambria Math')
        fonts.set(qn('w:hAnsi'), 'Cambria Math')
        pr.append(fonts)
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '24')
        pr.append(sz)
        run.insert(0, pr)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = Cm(0)
    para._p.append(node)
    para.add_run(f'    ({n})')
    para.paragraph_format.keep_with_next = True
    p(meaning)
    equations.append(dict(number=n, latex=tex, meaning=meaning))
    (FORM / f'eq_{n:02}.mml').write_bytes(etree.tostring(ml, encoding='utf-8', xml_declaration=True))
    prose.append('$$' + tex + '$$')


def table(head, data, widths=None, size=10.5):
    tb = doc.add_table(rows=1, cols=len(head))
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tb.autofit = False
    if widths:
        for col, w in zip(tb.columns, widths):
            col.width = Cm(w)
    for c, t in zip(tb.rows[0].cells, head):
        c.text = str(t)
    repeat = OxmlElement('w:tblHeader')
    tb.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in data:
        for c, t in zip(tb.add_row().cells, row):
            c.text = str(t)
    for i, row in enumerate(tb.rows):
        for k, c in enumerate(row.cells):
            if widths:
                c.width = Cm(widths[k])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = c._tc.get_or_add_tcPr()
            b = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                edge = OxmlElement('w:' + side)
                edge.set(qn('w:val'), 'single')
                edge.set(qn('w:sz'), '4')
                edge.set(qn('w:color'), '999999')
                b.append(edge)
            pr.append(b)
            if i == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'D9E2F3')
                pr.append(shade)
            for pa in c.paragraphs:
                pa.paragraph_format.line_spacing = 1.0
                pa.paragraph_format.space_after = Pt(1)
                pa.paragraph_format.space_before = Pt(1)
                pa.paragraph_format.first_line_indent = Cm(0)
                for run in pa.runs:
                    run.font.size = Pt(size)
                    run.bold = i == 0
                    run.font.name = 'Times New Roman'
                    run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.first_line_indent = Cm(0)
    prose.append('| ' + ' | '.join(head) + ' |\n' +
                 '\n'.join('| ' + ' | '.join(map(str, r)) + ' |' for r in data))


def fig(name, caption, width=15.5):
    pa = doc.add_paragraph()
    pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pa.paragraph_format.first_line_indent = Cm(0)
    pa.paragraph_format.keep_with_next = True
    pa.add_run().add_picture(str(F / (name + '.png')), width=Cm(width))
    doc.add_paragraph(caption, 'Caption')
    prose.append('![' + caption + '](../图表/' + name + '.png)')


# ==================== REPORT CONTENT ====================

doc.add_paragraph('问题3 可解释性多模态情感预测', 'Title')
doc.add_paragraph('——问题2矩阵补全方法的迁移验证与可解释性分析', 'Subtitle')

p('摘要：针对赛题问题3“三模态信息完整但决策依据难追溯”的场景，本报告将问题2建立的低秩矩阵补全情感预测模型（matrix_no_aug）迁移应用于问题3数据，在不重新训练模型的前提下，对附件2验证集和附件4测试集完成预测、模态归因与局部证据定位。模型在728条验证样本上取得Accuracy 0.6319、Macro-F1 0.5946、MAE 0.5965、Pearson 0.6390，略低于问题3专项模型（Accuracy 0.6484、Macro-F1 0.6299、MAE 0.5656），但解释忠实性实验表明积分梯度排序的关键位置仍能有效定位模型决策依据（20%位置替换时类别概率下降0.2751，显著高于注意力的0.0712和随机的0.0454）。附件4共20条样本完成预测与三模态Shapley贡献、词级积分梯度归因。结果表明，问题2方法可直接迁移至问题3完整模态场景，虽精度略低，但可解释性框架完整可用。', indent=False)

p('关键词：多模态情感识别；低秩矩阵补全；Shapley值；积分梯度；可解释性', indent=False)

# ---------- 1 问题重述 ----------
h('1 问题重述与方法迁移思路')
p('问题3要求在三模态信息相对完整的场景下，构建可解释性多模态情感预测模型，同时输出情感极性、情感强度，并以可量化、可复核方式说明判断依据，包括主要参考模态、模态作用程度和关键证据定位。问题2已建立一套面向局部模态缺失的低秩矩阵补全预测模型，本报告将该模型直接迁移至问题3数据，验证其在完整模态场景下的预测性能与可解释性。')
p('迁移策略如下：（1）直接复用问题2训练完成的matrix_no_aug模型权重与标准化参数，不重新训练；（2）输入使用官方aligned_50特征，其中text字段经核实与冻结BERT对text_bert的编码结果一致（最大绝对差5.8×10⁻⁵），因此无需重复编码；（3）在问题3的解释框架下，对该模型计算三模态精确Shapley值与词级积分梯度，验证解释忠实性。')

fig('图1_问题2模型结构', '图1 问题2低秩矩阵补全模型结构（迁移至问题3使用）')

# ---------- 2 模型结构 ----------
h('2 迁移模型结构')
p('问题2模型将文本（768维）、语音（74维）、视觉（35维）分别通过线性投影与双曲正切激活映射至32维隐空间，再经低秩矩阵补全模块融合三模态信息，最后通过注意力池化、均值池化与模态覆盖率拼接后送入预测头，输出三分类logits与连续情感强度。')

eq(r'z_{t,m}=O_{t,m}\tanh(W_m x_{t,m}+b_m),\quad z_{t,m}\in\mathbb{R}^{32}',
   '三模态特征投影至32维隐空间，O为观测掩码，仅在有效位置计算。')

eq(r'\min_U\ \sum_{t,d}w_{td}(Z_{td}-b_d^\top u_t)^2+\lambda_0\sum_t\|u_t\|_2^2+\lambda_T\sum_{t\sim t-1}\|u_t-u_{t-1}\|_2^2',
   '低秩矩阵补全目标：Z为拼接后的96维表示，B为共享字典（秩8），U为时间因子，λ₀=0.03为岭正则，λ_T=0.05为时间正则。')

eq(r'A_t=B^\top\operatorname{diag}(w_t)B+(\lambda_0+\lambda_T\deg t)I,\quad u_t^{k+1}=A_t^{-1}\left(b_t+\lambda_T\sum_{j\sim t}u_j^k\right)',
   '通过Cholesky分解求逆后进行20次块Jacobi迭代求解时间因子。已观测位置保留原始投影表示，仅缺失位置采用裁切至[-1,1]的重建值。')

eq(r'\widehat y=3\tanh(o_r),\quad \widehat c=\operatorname{argmax}\operatorname{softmax}(o_c)',
   '输出规则：情感强度限制在[-3,3]，极性由三分类logits的最大值决定（classification_projected规则）。')

# ---------- 3 可解释性方法 ----------
h('3 可解释性归因方法')
p('为使问题2模型满足问题3的可解释性要求，本报告采用精确Shapley值刻画模态整体贡献，采用积分梯度刻画局部特征贡献，并通过删除实验验证解释忠实性。')

h('3.1 三模态精确Shapley值', level=2)
p('对三个模态枚举全部8种组合，以标准化零向量（训练均值）为基线，保留样本原有有效位置掩码。分别对情感强度输出和预测类别概率计算Shapley值。')

eq(r'v(S)=f(X_S,0_{\overline S};M),\qquad S\subseteq\{T,A,V\}',
   '联盟价值函数：S为包含的模态集合，其余模态特征替换为零。')

eq(r'\phi_m=\sum_{S\in\mathcal S_m}\frac{|S|!(2-|S|)!}{3!}\left[v(S\cup\{m\})-v(S)\right]',
   '模态m的Shapley值，精确枚举3个模态的全部8种组合。')

eq(r'\eta_m=\frac{|\phi_m|}{\sum_j|\phi_j|},\qquad m^*=\operatorname{argmax}_m\eta_m',
   '模态作用程度η按预测类别概率的Shapley绝对值归一化，主要参考模态为绝对贡献最大者。')

h('3.2 积分梯度', level=2)
p('使用64点Gauss-Legendre求积计算输入特征对输出的积分梯度，定位局部关键位置。')

eq(r'\operatorname{IG}_{td}^{(m)}=\widetilde X_{td}^{(m)}\int_0^1 g_{td}^{(m)}(\alpha\widetilde X;M)\,d\alpha',
   'g为模型对输入坐标的偏导数，在α倍输入处取值。同一原文词的多个词元贡献聚合成词级柱状图。')

h('3.3 解释忠实性验证', level=2)
p('删除实验将有效位置中前r比例替换为训练均值（特征置零，掩码不变），比较积分梯度、注意力权重和均匀随机三种排序方式导致的预测类别概率下降。')

eq(r'D(r)=\frac1N\sum_i\left[p_{\widehat c_i}(X_i)-p_{\widehat c_i}(X_i^{\mathrm{replace}(r)})\right]',
   'D(r)为替换r比例位置后原预测类别概率的平均下降量，越大表示所删位置越影响判断。')

# ---------- 4 验证集结果 ----------
h('4 验证集结果与模型对比')

table(['模型', 'Accuracy', 'Macro-F1', 'MAE', 'Pearson'],
      [['问题2模型 (matrix_no_aug)', f"{met['accuracy']:.4f}", f"{met['macro_f1']:.4f}",
        f"{met['mae']:.4f}", f"{met['pearson']:.4f}"],
       ['问题3模型 (fusion_h32)', '0.6484', '0.6299', '0.5656', '0.6851'],
       ['差值 (Q2−Q3)', f"{met['accuracy']-0.6484:+.4f}", f"{met['macro_f1']-0.6299:+.4f}",
        f"{met['mae']-0.5656:+.4f}", f"{met['pearson']-0.6851:+.4f}"]],
      [4.5, 2.5, 2.5, 2.5, 2.5])

p('问题2模型在验证集上的Accuracy较问题3模型低1.65个百分点，Macro-F1低0.0353，MAE高0.0309，Pearson低0.0461。差距主要来源于两模型结构差异：问题3模型使用SiLU激活、深度卷积与非线性融合，而问题2模型使用Tanh激活、矩阵补全与线性投影头。问题2模型未针对完整模态场景优化，但其预测性能仍处于可用范围。')

fig('图2_问题2与问题3模型验证对比', '图2 问题2迁移模型与问题3专项模型验证集指标对比')

fig('图3_混淆矩阵与回归散点', f'图3 验证集混淆矩阵与回归散点（MAE={met["mae"]:.4f}，Pearson={met["pearson"]:.4f}）')

# ---------- 5 模态作用差异 ----------
h('5 三模态作用差异')
counts = audit['valid']['principal_modality_counts']
total = sum(counts)
table(['主要参考模态', '样本数', '占比'],
      [['文本', counts[0], f'{counts[0]/total:.2%}'],
       ['语音', counts[1], f'{counts[1]/total:.2%}'],
       ['视觉', counts[2], f'{counts[2]/total:.2%}']],
      [4, 3, 3])

p(f'验证集中主要参考文本、语音、视觉的样本分别为{counts[0]}、{counts[1]}、{counts[2]}条，占比{counts[0]/total:.2%}、{counts[1]/total:.2%}、{counts[2]/total:.2%}。与问题3模型（文本84.34%、语音6.04%、视觉9.62%）的分布高度一致，说明两模型在模态依赖模式上相似，文本在多数样本中占主导。')

fig('图4_验证集模态作用差异', '图4 验证集各模态作用程度分布及主要参考模态数量')

# ---------- 6 解释忠实性 ----------
h('6 解释忠实性验证')
faith_20 = {r['method']: float(r['probability_drop_mean']) for r in faith
            if float(r['rate']) == .2 and int(r['repeat']) == 0}
table(['排序方法', '20%替换时概率下降'],
      [['积分梯度关键位置', f'{faith_20["integrated_gradients"]:.4f}'],
       ['注意力权重位置', f'{faith_20["attention"]:.4f}'],
       ['均匀伪随机位置', f'{faith_20["hashed_uniform"]:.4f}']],
      [5, 5])

p(f'20%位置替换时，积分梯度、注意力和均匀位置的平均概率下降分别为{faith_20["integrated_gradients"]:.4f}、{faith_20["attention"]:.4f}和{faith_20["hashed_uniform"]:.4f}。积分梯度排序的关键位置对模型输出的影响远大于注意力和随机排序，验证了积分梯度归因的忠实性。与问题3模型（IG=0.2953、Attn=0.1535、Rand=0.0427）相比，问题2模型的IG忠实性略低但仍显著有效。')

fig('图5_解释忠实性验证', '图5 相同替换比例下原预测类别概率下降（越大表示所删位置越影响判断）')

# ---------- 7 错误分析 ----------
h('7 错误分布与归因')
fig('图7_验证集错误分组', '图7 按真实强度分组的错误分布')
p('从错误分组看，强情感样本（|y|>1.5）的分类准确率较高但强度回归存在向中间收缩的现象，MAE超过1.0；中性与弱情感样本在类别边界附近容易混淆。最大误差样本包括比较对象混淆（如"convenient"与"stupid boots"出现在同一句中）、隐含评价（无明显情感词但态度负向）以及50位置截断导致关键信息丢失等情况，与问题3模型的错误模式一致。')

# ---------- 8 附件4预测 ----------
h('8 附件4全量预测与解释')
polarity_map = {'Negative': '负向', 'Neutral': '中性', 'Positive': '正向'}
mod_map = {'text': '文本', 'audio': '语音', 'vision': '视觉'}
pred_table = []
for x in pred:
    pred_table.append([x['sample_id'], polarity_map[x['polarity']], f"{float(x['intensity']):+.3f}",
                       mod_map[x['principal_modality']],
                       f"{float(x['text_effect'])*100:.1f}",
                       f"{float(x['audio_effect'])*100:.1f}",
                       f"{float(x['vision_effect'])*100:.1f}"])
table(['编号', '极性', '强度', '主要模态', '文本%', '语音%', '视觉%'],
      pred_table, [1.2, 1.5, 2, 2, 2, 2, 2], size=9.5)

neg = sum(1 for x in pred if x['polarity'] == 'Negative')
neu = sum(1 for x in pred if x['polarity'] == 'Neutral')
pos = sum(1 for x in pred if x['polarity'] == 'Positive')
p(f'20条样本中预测负向{neg}条、中性{neu}条、正向{pos}条。此为预测分布，不代表真实类别。主参考模态为文本的样本最多，少数样本以语音或视觉为主。')

fig('图6_附件4全量预测与解释', '图8 附件4全部20条样本的模态作用占比和强度预测')

# ---------- 9 典型样本 ----------
h('9 典型样本解释卡')
a4_dir = Path(r'<project_root>\附件4-可解释专项视频样本与特征文件\附件4-可解释专项视频样本与特征文件\对齐版本')
a4_raw = {}
for pf in sorted(a4_dir.glob('*.pkl')):
    import pickle
    with pf.open('rb') as f:
        d = pickle.load(f)
    a4_raw[d['id']] = d['raw_text']

for sid in ['01', '02', '03', '04', '05']:
    row = next(x for x in pred if x['sample_id'] == sid)
    raw = a4_raw.get(sid, '')
    h(f'9.{int(sid)} 样本{sid}（{mod_map[row["principal_modality"]]}主导）', level=2)
    p(f'原始文本：{raw}', indent=False)
    fig(f'样本{sid}_解释卡', f'图9 样本{sid}的预测、模态贡献与词级积分梯度', width=15.5)
    p(f'样本{sid}预测{polarity_map[row["polarity"]]}、强度{float(row["intensity"]):+.3f}，主要参考{mod_map[row["principal_modality"]]}（占比{float(row[row["principal_modality"]+"_effect"])*100:.1f}%），类别概率{float(row["predicted_class_probability"]):.3f}。')

# ---------- 10 结论 ----------
h('10 结论')
p('本报告将问题2建立的低秩矩阵补全模型直接迁移至问题3完整模态场景，得出以下结论：')
p('（1）问题2模型无需重新训练即可在问题3数据上运行，验证集Accuracy 0.6319、Macro-F1 0.5946、MAE 0.5965，较问题3专项模型有小幅差距，但预测性能仍处于合理范围。')
p('（2）可解释性框架完整可用。三模态精确Shapley值（残差2.38×10⁻⁷）和积分梯度（残差4.41×10⁻⁶）均满足完整性约束，解释忠实性实验验证了积分梯度归因的有效性。')
p('（3）两模型在模态依赖模式上高度一致，文本均占主导（约84%），说明问题2方法捕捉到了与问题3模型相似的多模态融合规律。')
p('（4）问题2模型的矩阵补全模块在完整模态场景下虽未发挥缺失修复作用，但其低秩融合结构仍能有效整合三模态信息，证明了该方法的跨场景迁移能力。')

h('参考文献')
p('[1] Lundberg S M, Lee S I. A Unified Approach to Interpreting Model Predictions. NeurIPS, 2017.', indent=False)
p('[2] Sundararajan M, Taly A, Yan Q. Axiomatic Attribution for Deep Networks. ICML, 2017.', indent=False)
p('[3] 赛题原文《复杂场景下多模态情感识别的数学建模与算法设计.docx》。', indent=False)

# ---- Formula appendix ----
doc.add_page_break()
h('附录 公式LaTeX源码')
for item in equations:
    doc.add_heading(f'公式 {item["number"]}', level=2)
    pa = doc.add_paragraph()
    pa.paragraph_format.first_line_indent = Cm(0)
    run = pa.add_run(item['latex'])
    run.font.name = 'Consolas'
    run.font.size = Pt(10)
    p(item['meaning'])

tex = '\\documentclass{article}\n\\usepackage{amsmath,amssymb}\n\\begin{document}\n' + \
      '\n'.join('\\begin{equation}\n' + e['latex'] + '\n\\end{equation}' for e in equations) + \
      '\n\\end{document}\n'
(FORM / '全部公式.tex').write_text(tex, encoding='utf-8')
(FORM / '全部公式.json').write_text(json.dumps(equations, ensure_ascii=False, indent=2), encoding='utf-8')
(FORM / '全部公式.txt').write_text('\n\n'.join(f'({e["number"]}) {e["latex"]}' for e in equations), encoding='utf-8')

# ---- Save ----
docx_path = OUT / '问题3_问题2方法迁移验证报告.docx'
doc.save(docx_path)
(OUT / '问题3_问题2方法迁移验证报告.md').write_text('\n\n'.join(prose), encoding='utf-8')
(OUT / '公式格式说明.txt').write_text(
    'Word正文为可编辑OMML数学对象（Cambria Math字体），非图片。'
    'MathType原生OLE转换请在Word中使用：MathType > Convert Equations > '
    'Word OMML (2007 and later) > Whole document > MathType OLE。'
    '正文字体为Times New Roman（西文）与宋体（中文）。',
    encoding='utf-8')

print(json.dumps({'docx': str(docx_path), 'equations': len(equations),
                  'figures': len(list(F.glob('*.png')))}, ensure_ascii=False))
